import os
import shutil
import logging
from pathlib import Path
from datetime import datetime
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from modules.config import (
    EXCEL_FILE_PATH,
    EXCEL_COLUMNS,
    COLUMN_WIDTHS,
    EXCEL_HEADER_BG,
    EXCEL_HEADER_FG,
    EXCEL_ZEBRA_BG,
    EXCEL_BORDER_COLOR,
    DATA_DIR
)
from modules.utils import format_image_id

logger = logging.getLogger("excel_manager")

class ExcelManager:
    def __init__(self, file_path: Path | str = EXCEL_FILE_PATH):
        self.file_path = Path(file_path)
        self._ensure_workbook_exists()

    def _ensure_workbook_exists(self) -> None:
        """Initializes the Excel workbook with styled headers if it does not exist."""
        self.file_path.parent.mkdir(parents=True, exist_ok=True)
        if not self.file_path.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = "EventDhara Metadata"
            self._apply_header_styles(ws)
            wb.save(self.file_path)
            # Also keep a copy in data directory
            self._sync_backup()

    def _sync_backup(self) -> None:
        """Saves a backup copy into the data directory."""
        try:
            backup_path = DATA_DIR / "EventDhara_Metadata_backup.xlsx"
            if self.file_path.exists():
                shutil.copy2(self.file_path, backup_path)
        except Exception as e:
            logger.warning(f"Could not create backup copy: {e}")

    def _apply_header_styles(self, ws) -> None:
        """Applies executive styling, font, fill, and borders to the header row."""
        header_font = Font(name="Segoe UI", size=11, bold=True, color=EXCEL_HEADER_FG)
        header_fill = PatternFill(
            start_color=EXCEL_HEADER_BG,
            end_color=EXCEL_HEADER_BG,
            fill_type="solid"
        )
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin = Side(border_style="thin", color=EXCEL_BORDER_COLOR)
        border = Border(top=thin, left=thin, right=thin, bottom=thin)

        ws.row_dimensions[1].height = 32

        for col_idx, header_title in enumerate(EXCEL_COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header_title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border

        ws.freeze_panes = "A2"
        self._apply_column_widths(ws)

    def _apply_column_widths(self, ws) -> None:
        """Applies configured readable column widths."""
        for col_idx, width in COLUMN_WIDTHS.items():
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

    def get_existing_records(self) -> list[dict]:
        """Reads all records from the Excel file."""
        if not self.file_path.exists():
            return []

        wb = load_workbook(self.file_path, data_only=True)
        ws = wb.active
        records = []

        headers = [ws.cell(row=1, column=c).value for c in range(1, len(EXCEL_COLUMNS) + 1)]

        for r in range(2, ws.max_row + 1):
            row_dict = {}
            has_data = False
            for c_idx, h in enumerate(headers, 1):
                val = ws.cell(row=r, column=c_idx).value
                if val is not None and str(val).strip():
                    has_data = True
                row_dict[h] = val

            if has_data and row_dict.get("S.No") is not None:
                records.append(row_dict)

        return records

    def check_duplicate(self, url: str | None = None, image_hash: str | None = None) -> dict:
        """
        Checks if a URL or Image Hash already exists in the catalog.
        
        Returns:
            {
                "is_duplicate": bool,
                "reason": "url" | "image_hash" | None,
                "row": int | None,
                "existing_record": dict | None
            }
        """
        records = self.get_existing_records()
        clean_url = (url or "").strip().lower()
        clean_hash = (image_hash or "").strip().lower()

        for idx, rec in enumerate(records, 2):
            rec_url = str(rec.get("Original Pinterest URL") or "").strip().lower()
            rec_hash = str(rec.get("Image Hash (MD5)") or "").strip().lower()

            # URL check (ignore generic labels like 'Manual Upload')
            if clean_url and "manual upload" not in clean_url and rec_url == clean_url:
                return {
                    "is_duplicate": True,
                    "reason": "url",
                    "row": idx,
                    "existing_record": rec
                }

            # Image Hash check
            if clean_hash and rec_hash == clean_hash:
                return {
                    "is_duplicate": True,
                    "reason": "image_hash",
                    "row": idx,
                    "existing_record": rec
                }

        return {
            "is_duplicate": False,
            "reason": None,
            "row": None,
            "existing_record": None
        }

    def append_record(self, record_data: dict, force: bool = False) -> dict:
        """
        Appends a new metadata record with sequential serial number and styling.
        Prevents duplicate entry unless force=True.
        
        Returns:
            {
                "success": bool,
                "sno": int,
                "image_id": str,
                "row": int,
                "message": str
            }
        """
        url = record_data.get("Original Pinterest URL", "")
        img_hash = record_data.get("Image Hash (MD5)", "")

        if not force:
            dup_check = self.check_duplicate(url=url, image_hash=img_hash)
            if dup_check["is_duplicate"]:
                existing = dup_check["existing_record"]
                return {
                    "success": False,
                    "is_duplicate": True,
                    "sno": existing.get("S.No"),
                    "image_id": existing.get("Image ID"),
                    "row": dup_check["row"],
                    "message": f"Duplicate detected ({dup_check['reason']}): Already recorded as S.No #{existing.get('S.No')} (ID: {existing.get('Image ID')})"
                }

        wb = load_workbook(self.file_path)
        ws = wb.active

        # Determine next serial number
        max_sno = 0
        for r in range(2, ws.max_row + 1):
            s_val = ws.cell(row=r, column=1).value
            if s_val is not None:
                try:
                    s_int = int(s_val)
                    if s_int > max_sno:
                        max_sno = s_int
                except (ValueError, TypeError):
                    pass

        next_sno = max_sno + 1
        target_row = ws.max_row + 1 if ws.cell(row=1, column=1).value else 2

        # Assign standardized ID and serial number
        image_id = record_data.get("Image ID") or format_image_id(next_sno)
        record_data["S.No"] = next_sno
        record_data["Image ID"] = image_id
        if "Analyzed At" not in record_data or not record_data["Analyzed At"]:
            record_data["Analyzed At"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # Styling
        data_font = Font(name="Segoe UI", size=10)
        thin = Side(border_style="thin", color=EXCEL_BORDER_COLOR)
        border = Border(top=thin, left=thin, right=thin, bottom=thin)

        is_even = (target_row % 2 == 0)
        row_fill = PatternFill(start_color=EXCEL_ZEBRA_BG, end_color=EXCEL_ZEBRA_BG, fill_type="solid") if is_even else None

        ws.row_dimensions[target_row].height = 42

        for col_idx, col_name in enumerate(EXCEL_COLUMNS, 1):
            val = record_data.get(col_name, "")
            cell = ws.cell(row=target_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = border
            if row_fill:
                cell.fill = row_fill

            # Alignment rules
            if col_name in ["S.No", "Image ID", "Analyzed At", "Event Category"]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_name in ["Short Description", "Long Description", "Image Inclusions", "Keywords and Search Tags"]:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

        # Refresh auto-filter and dimensions
        ws.auto_filter.ref = ws.dimensions
        self._apply_column_widths(ws)

        wb.save(self.file_path)
        self._sync_backup()

        return {
            "success": True,
            "is_duplicate": False,
            "sno": next_sno,
            "image_id": image_id,
            "row": target_row,
            "message": f"Successfully recorded as S.No #{next_sno} ({image_id})"
        }

    def get_as_dataframe(self) -> pd.DataFrame:
        """Returns the master Excel records as a clean Pandas DataFrame."""
        records = self.get_existing_records()
        if not records:
            return pd.DataFrame(columns=EXCEL_COLUMNS)
        df = pd.DataFrame(records)
        # Ensure column ordering matches
        cols_present = [c for c in EXCEL_COLUMNS if c in df.columns]
        return df[cols_present]
